from __future__ import annotations

import csv
import io
import json
import logging
import os
import queue
import threading
import time
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import desc, select
from sqlalchemy.engine import Engine

from billetvision import config
from billetvision.db import get_engine, insert_inspection, inspections_table

logger = logging.getLogger(__name__)

LOG_COLUMNS: tuple[str, ...] = (
    "id",
    "timestamp",
    "billet_id",
    "id_source",
    "ocr_text",
    "ocr_confidence",
    "qr_text",
    "id_mismatch",
    "length_mm",
    "width_mm",
    "height_mm",
    "length_dev_mm",
    "width_dev_mm",
    "width_variation_pct",
    "taper_mm",
    "camber_mm",
    "defects",
    "status",
    "auto_status",
    "reasons",
    "spec_id",
    "n_samples",
    "latency_ms",
    "override_note",
    "snapshot",
)


def _format_cell_value(key: str, val: Any) -> Any:
    """Format row fields for CSV and XLSX representation."""
    if val is None:
        return ""
    if key in (
        "length_mm",
        "width_mm",
        "height_mm",
        "length_dev_mm",
        "width_dev_mm",
        "width_variation_pct",
        "taper_mm",
        "camber_mm",
        "latency_ms",
    ):
        try:
            return round(float(val), 2)
        except (ValueError, TypeError):
            return val
    if key == "ocr_confidence":
        try:
            return round(float(val), 2)
        except (ValueError, TypeError):
            return ""
    if key in ("defects", "reasons"):
        if isinstance(val, (list, tuple)):
            return "; ".join(str(x) for x in val)
        if isinstance(val, str):
            try:
                parsed = json.loads(val)
                if isinstance(parsed, list):
                    return "; ".join(str(x) for x in parsed)
            except Exception:
                pass
            return val
    if key == "id_mismatch":
        return bool(val)
    return str(val)


def row_to_export_dict(row: dict[str, Any]) -> dict[str, Any]:
    """Map DB or internal row keys to the standard export column schema."""
    res: dict[str, Any] = {}
    res["id"] = row.get("id", "")
    res["timestamp"] = row.get("ts") or row.get("timestamp", "")
    res["billet_id"] = row.get("billet_id") or ""
    res["id_source"] = row.get("id_source") or ""
    res["ocr_text"] = row.get("ocr_text") or ""
    res["ocr_confidence"] = _format_cell_value("ocr_confidence", row.get("ocr_confidence"))
    res["qr_text"] = row.get("qr_text") or ""
    res["id_mismatch"] = _format_cell_value("id_mismatch", row.get("id_mismatch", False))
    res["length_mm"] = _format_cell_value("length_mm", row.get("length_mm"))
    res["width_mm"] = _format_cell_value("width_mm", row.get("width_mm"))
    res["height_mm"] = _format_cell_value("height_mm", row.get("height_mm"))
    res["length_dev_mm"] = _format_cell_value("length_dev_mm", row.get("length_dev_mm"))
    res["width_dev_mm"] = _format_cell_value("width_dev_mm", row.get("width_dev_mm"))

    var_pct = row.get("width_variation_pct") if "width_variation_pct" in row else row.get("cs_var_pct")
    res["width_variation_pct"] = _format_cell_value("width_variation_pct", var_pct)

    res["taper_mm"] = _format_cell_value("taper_mm", row.get("taper_mm"))
    res["camber_mm"] = _format_cell_value("camber_mm", row.get("camber_mm"))
    res["defects"] = _format_cell_value("defects", row.get("defects"))
    res["status"] = row.get("status") or ""
    res["auto_status"] = row.get("auto_status") or ""
    res["reasons"] = _format_cell_value("reasons", row.get("reasons"))
    res["spec_id"] = row.get("spec_id") or ""
    res["n_samples"] = row.get("n_samples", 0)
    res["latency_ms"] = _format_cell_value("latency_ms", row.get("latency_ms"))

    override_val = row.get("override")
    note = ""
    if isinstance(override_val, dict):
        note = override_val.get("note", "")
    elif isinstance(override_val, str) and override_val:
        try:
            parsed = json.loads(override_val)
            if isinstance(parsed, dict):
                note = parsed.get("note", "")
        except Exception:
            note = override_val
    res["override_note"] = row.get("override_note") or note

    res["snapshot"] = row.get("snapshot_path") or row.get("snapshot_url") or ""
    return res


def build_csv(rows: Sequence[dict[str, Any]]) -> str:
    """Generate RFC 4180 CSV export as string."""
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=list(LOG_COLUMNS))
    writer.writeheader()
    for row in rows:
        writer.writerow(row_to_export_dict(row))
    return output.getvalue()


def build_json(rows: Sequence[dict[str, Any]]) -> str:
    """Generate JSON export as string."""
    exported = [row_to_export_dict(r) for r in rows]
    return json.dumps(exported, indent=2)


def build_workbook(
    rows: Sequence[dict[str, Any]],
    active_spec: str = "FLAT-500x100",
) -> bytes:
    """Build styled Excel workbook per SPEC §14 & 03 §5."""
    wb = openpyxl.Workbook()

    # Sheet 1: Inspections
    ws = wb.active
    ws.title = "Inspections"

    # Status color fills
    status_fills = {
        "Pass": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "Rework": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "Fail": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "Review": PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"),
    }
    status_fonts = {
        "Pass": Font(color="006100", bold=True),
        "Rework": Font(color="9C6500", bold=True),
        "Fail": Font(color="9C0006", bold=True),
        "Review": Font(color="1F497D", bold=True),
    }

    # Headers
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F3542", end_color="2F3542", fill_type="solid")

    ws.append(list(LOG_COLUMNS))
    for col_idx in range(1, len(LOG_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    status_col_idx = LOG_COLUMNS.index("status") + 1

    for row_idx, r in enumerate(rows, start=2):
        exp = row_to_export_dict(r)
        row_vals = [exp[col] for col in LOG_COLUMNS]
        ws.append(row_vals)

        # Style status cell
        status_val = str(exp["status"])
        if status_val in status_fills:
            c = ws.cell(row=row_idx, column=status_col_idx)
            c.fill = status_fills[status_val]
            c.font = status_fonts[status_val]
            c.alignment = Alignment(horizontal="center")

    # Freeze top row
    ws.freeze_panes = "A2"

    # Enable auto filter
    if rows:
        last_col_letter = get_column_letter(len(LOG_COLUMNS))
        ws.auto_filter.ref = f"A1:{last_col_letter}{len(rows) + 1}"

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    # Sheet 2: Summary
    ws_summary = wb.create_sheet(title="Summary")
    title_font = Font(name="Calibri", size=14, bold=True)
    section_font = Font(name="Calibri", size=11, bold=True)

    ws_summary["A1"] = "BilletVision Inspection Log Summary"
    ws_summary["A1"].font = title_font

    ws_summary["A3"] = "Generated At:"
    ws_summary["B3"] = time.strftime("%Y-%m-%d %H:%M:%S UTC", time.gmtime())

    ws_summary["A4"] = "Active Spec:"
    ws_summary["B4"] = active_spec

    ws_summary["A5"] = "Total Inspections:"
    ws_summary["B5"] = len(rows)

    by_status: dict[str, int] = {
        "Pass": 0,
        "Rework": 0,
        "Fail": 0,
        "Review": 0,
    }
    for r in rows:
        st = r.get("status")
        if st in by_status:
            by_status[st] += 1

    ws_summary["A7"] = "Status"
    ws_summary["B7"] = "Count"
    ws_summary["C7"] = "Share"
    for col_l in ("A7", "B7", "C7"):
        ws_summary[col_l].font = section_font

    tot = max(len(rows), 1)
    s_row = 8
    for st_name, count in by_status.items():
        ws_summary.cell(row=s_row, column=1, value=st_name)
        ws_summary.cell(row=s_row, column=2, value=count)
        share_pct = (count / tot) * 100.0 if rows else 0.0
        ws_summary.cell(row=s_row, column=3, value=f"{share_pct:.1f}%")
        s_row += 1

    ws_summary.column_dimensions["A"].width = 22
    ws_summary.column_dimensions["B"].width = 16
    ws_summary.column_dimensions["C"].width = 14

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


class InspectionLogger:
    """Single writer thread handling SQLite, CSV append, and XLSX atomic replacement."""

    def __init__(
        self,
        log_dir: str | Path | None = None,
        db_engine: Engine | None = None,
        queue_size: int = config.JOB_QUEUE_SIZE,
        xlsx_sync_interval: int = 5,
    ) -> None:
        self.log_dir = Path(log_dir or config.LOG_DIR).resolve()
        self.log_dir.mkdir(parents=True, exist_ok=True)
        self.csv_path = self.log_dir / "inspections.csv"
        self.xlsx_path = self.log_dir / "inspections.xlsx"

        self.db_engine = db_engine or get_engine()
        self.queue: queue.Queue[dict[str, Any]] = queue.Queue(maxsize=queue_size)
        self._csv_pending: list[dict[str, Any]] = []

        self._last_successful_write: float | None = None
        self._blocked: bool = False
        self._xlsx_counter: int = 0
        self._xlsx_dirty: bool = False
        self._last_xlsx_time: float = 0.0
        self._xlsx_sync_interval: int = xlsx_sync_interval

        self._stop_event = threading.Event()
        self._lock = threading.Lock()
        self._worker_thread = threading.Thread(target=self._run_loop, name="logging_io_worker", daemon=True)
        self._worker_thread.start()

    @property
    def queue_length(self) -> int:
        return self.queue.qsize()

    @property
    def last_successful_write_time(self) -> float | None:
        return self._last_successful_write

    @property
    def blocked(self) -> bool:
        return self._blocked

    @property
    def pending_count(self) -> int:
        with self._lock:
            return len(self._csv_pending)

    def log_inspection(
        self,
        record: dict[str, Any],
        block: bool = False,
        timeout: float = 0.1,
    ) -> bool:
        """Enqueue an inspection record without blocking the vision loop."""
        try:
            self.queue.put(record, block=block, timeout=timeout if block else None)
            return True
        except queue.Full:
            logger.warning("InspectionLogger queue full; dropping oldest job to prevent producer stalls")
            try:
                self.queue.get_nowait()
            except queue.Empty:
                pass
            try:
                self.queue.put_nowait(record)
                return True
            except queue.Full:
                return False

    def _append_to_csv(self, records: Sequence[dict[str, Any]]) -> None:
        """Append rows to CSV with fsync; raises on PermissionError / OSError."""
        if not records:
            return
        needs_header = not self.csv_path.exists() or self.csv_path.stat().st_size == 0
        with open(self.csv_path, "a", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(LOG_COLUMNS))
            if needs_header:
                writer.writeheader()
            for r in records:
                writer.writerow(row_to_export_dict(r))
            f.flush()
            os.fsync(f.fileno())

    def _regenerate_xlsx(self) -> None:
        """Regenerate XLSX atomically from SQLite; fall back to timestamped file if locked."""
        try:
            with self.db_engine.connect() as conn:
                stmt = (
                    select(inspections_table)
                    .order_by(desc(inspections_table.c.id))
                    .limit(config.EXCEL_MAX_ROWS)
                )
                rows = [dict(r) for r in conn.execute(stmt).mappings().all()]
            # Reverse so chronological order in sheet
            rows.reverse()

            wb_bytes = build_workbook(rows)
            tmp_path = self.log_dir / ".inspections.tmp.xlsx"
            tmp_path.write_bytes(wb_bytes)

            try:
                os.replace(tmp_path, self.xlsx_path)
                with self._lock:
                    self._xlsx_dirty = False
                    self._xlsx_counter = 0
                    self._last_xlsx_time = time.time()
                    if not self._csv_pending:
                        self._blocked = False
            except (PermissionError, OSError) as pe:
                logger.warning(
                    f"inspections.xlsx locked by another process ({pe}); writing timestamped fallback"
                )
                ts_name = f"billet_log_{int(time.time())}.xlsx"
                fallback_path = self.log_dir / ts_name
                try:
                    os.replace(tmp_path, fallback_path)
                except Exception:
                    fallback_path.write_bytes(wb_bytes)
                with self._lock:
                    self._blocked = True
        except Exception as e:
            logger.error(f"Error regenerating XLSX: {e}", exc_info=True)
            with self._lock:
                self._blocked = True

    def _run_loop(self) -> None:
        last_csv_retry = 0.0

        while not self._stop_event.is_set():
            try:
                record = self.queue.get(timeout=0.2)
            except queue.Empty:
                record = None

            now = time.time()

            # Handle new record
            if record is not None:
                # 1. SQLite first (source of truth)
                try:
                    ins_id = insert_inspection(record, engine=self.db_engine)
                    record["id"] = ins_id
                except Exception as e:
                    logger.error(f"Database write failure in logger: {e}", exc_info=True)

                # 2. Append to CSV
                with self._lock:
                    self._csv_pending.append(record)
                    to_write = list(self._csv_pending)

                try:
                    self._append_to_csv(to_write)
                    with self._lock:
                        self._csv_pending.clear()
                        self._last_successful_write = time.time()
                        self._blocked = False
                except (PermissionError, OSError) as e:
                    logger.warning(f"CSV append failed ({e}); buffered in memory")
                    with self._lock:
                        self._blocked = True

                with self._lock:
                    self._xlsx_dirty = True
                    self._xlsx_counter += 1

                self.queue.task_done()

            # Retry CSV pending if any
            with self._lock:
                pending_exist = bool(self._csv_pending)

            if pending_exist and (now - last_csv_retry >= 1.0):
                last_csv_retry = now
                with self._lock:
                    to_retry = list(self._csv_pending)
                try:
                    self._append_to_csv(to_retry)
                    with self._lock:
                        self._csv_pending.clear()
                        self._last_successful_write = time.time()
                        self._blocked = False
                    logger.info("Successfully flushed pending CSV records")
                except (PermissionError, OSError):
                    pass

            # Regenerate XLSX if threshold met
            with self._lock:
                dirty = self._xlsx_dirty
                cnt = self._xlsx_counter
                last_x = self._last_xlsx_time

            if dirty and (cnt >= self._xlsx_sync_interval or (now - last_x >= config.EXCEL_DEBOUNCE_S)):
                self._regenerate_xlsx()

    def flush_sync(self, timeout: float = 5.0) -> None:
        """Block until the queue is drained and XLSX is refreshed."""
        t0 = time.time()
        while not self.queue.empty() and (time.time() - t0 < timeout):
            time.sleep(0.05)
        # Flush pending CSV if possible
        with self._lock:
            pending = list(self._csv_pending)
        if pending:
            try:
                self._append_to_csv(pending)
                with self._lock:
                    self._csv_pending.clear()
            except Exception:
                pass
        self._regenerate_xlsx()

    def stop(self, timeout: float = 3.0) -> None:
        """Signal thread shutdown and join."""
        self.flush_sync(timeout=1.0)
        self._stop_event.set()
        self._worker_thread.join(timeout=timeout)
